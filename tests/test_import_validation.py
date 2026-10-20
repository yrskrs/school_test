"""Malformed user files must not become server errors or partial tests."""
import io
import json
from pathlib import Path
import sys
import unittest
import zipfile

_ROOT = Path(__file__).resolve().parents[1]
_TESTS_DIR = Path(__file__).resolve().parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))
if str(_TESTS_DIR) not in sys.path:
    sys.path.insert(0, str(_TESTS_DIR))

import test_flows as fixtures
from app import models


class ImportValidation(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = fixtures.Flows.asyncSetUp
    asyncTearDown = fixtures.Flows.asyncTearDown

    async def upload(self, path, name, content, field='file'):
        boundary = 'synthetic-import-validation'
        payload = (f'--{boundary}\r\nContent-Disposition: form-data; name="{field}"; filename="{name}"\r\n'
                   'Content-Type: application/octet-stream\r\n\r\n').encode() + content + f'\r\n--{boundary}--\r\n'.encode()
        return await fixtures.http('POST', path, payload, self.teacher_cookie,
            content_type='multipart/form-data; boundary=' + boundary)

    async def test_broken_docx_returns_400_without_creating_test(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('unrelated.txt', 'synthetic non-Word archive')
        before = self.db.query(models.Test).count()
        for content in (b'not a zip file', buffer.getvalue()):
            result = await self.upload('/teacher/tests/import/docx', 'broken.docx', content)
            self.assertEqual(result[0], 400, result[1])
            self.assertIn('DOCX', result[1])
        self.assertEqual(self.db.query(models.Test).count(), before)

    async def test_json_shape_and_encoding_errors_do_not_create_partial_tests(self):
        before = self.db.query(models.Test).count()
        for data in ([], None, {'title': 'Synthetic', 'questions': {}},
                     {'title': 'Synthetic', 'questions': [None]},
                     {'title': 'Synthetic', 'questions': [{'options': [None]}]},
                     {'title': 'Synthetic', 'questions': [{'points': -1}]},
                     {'title': 'Synthetic', 'time_limit_minutes': {}},
                     {'title': 'Synthetic', 'max_grade': 0},
                     {'title': 'Synthetic', 'subject': []},
                     {'title': 'Synthetic', 'questions': [{'image_url': {'bad': True}}]},
                     {'title': 'Synthetic', 'embedded_images': {'test.png': {}}}):
            with self.subTest(data=data):
                result = await fixtures.http('POST', '/teacher/tests/import-json',
                    {'json_text': json.dumps(data)}, self.teacher_cookie, form=True)
                self.assertEqual(result[0], 400, result[1][:200])
        result = await self.upload('/teacher/tests/import-json', 'bad.json', b'\xff\xfe', 'json_file')
        self.assertEqual(result[0], 400)
        self.assertIn('UTF-8', result[1])
        self.assertEqual(self.db.query(models.Test).count(), before)

    async def test_uppercase_docx_and_utf8_bom_json_import_successfully(self):
        from docx import Document
        document = Document()
        document.add_paragraph('1. Synthetic question?')
        document.add_paragraph().add_run('А) Synthetic correct').bold = True
        document.add_paragraph('Б) Synthetic other')
        buffer = io.BytesIO()
        document.save(buffer)
        result = await self.upload('/teacher/tests/import/docx', 'synthetic.DOCX', buffer.getvalue())
        self.assertEqual(result[0], 200, result[1])
        data = json.loads(result[1])
        imported = self.db.get(models.Test, data['test_id'])
        self.assertEqual(len(imported.questions), 1)
        self.assertTrue(any(option.is_correct for option in imported.questions[0].options))
        result = await self.upload('/teacher/tests/import-json', 'synthetic.json',
            b'\xef\xbb\xbf' + json.dumps({'title': 'Synthetic BOM test', 'questions': []}).encode(), 'json_file')
        self.assertEqual(result[0], 303, result[1][:200])


if __name__ == '__main__':
    unittest.main()
