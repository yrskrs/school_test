from __future__ import annotations

import base64
import csv
import io
import json
import os
import re
import uuid
import xml.etree.ElementTree as ET

try:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
except ImportError:
    Document = None
    WD_COLOR_INDEX = None

from sqlalchemy.orm import Session

from app import crud, models
from app.services.test_file_service import (
    generate_new_folder_name,
    move_temp_images_to_test,
    save_test_locally,
)

# ---------------------------------------------------------------------------
# JSON Import
# ---------------------------------------------------------------------------

def import_test_from_json(db: Session, teacher_id: int, json_data: str) -> models.Test:
    """
    Імпортує тест із JSON-рядка.
    Очікувана структура: { title, subject, class_name, description,
    time_limit_minutes, shuffle_questions, shuffle_options,
    show_result_after_finish, show_correct_answers, is_formative,
    questions: [{ question_text, question_type, points, topic, difficulty,
                  explanation, order_index,
                  options: [{ option_text, is_correct, order_index }] }] }
    """
    try:
        raw = json.loads(json_data)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Невірний JSON: {exc}")

    allowed_test_fields = {
        "title", "subject", "class_name", "description",
        "time_limit_minutes", "shuffle_questions", "shuffle_options",
        "show_result_after_finish", "show_correct_answers", "is_formative",
    }
    test_data = {k: v for k, v in raw.items() if k in allowed_test_fields}
    if "title" not in test_data:
        raise ValueError("Поле 'title' є обов'язковим")

    questions_raw = raw.get("questions", [])
    
    temp_session_id = None
    url_mapping = {}
    embedded_images = raw.get("embedded_images")
    if embedded_images and isinstance(embedded_images, dict):
        temp_session_id = f"import_json_{uuid.uuid4().hex}"
        for filename, b64 in embedded_images.items():
            ext = "." + filename.split(".")[-1] if "." in filename else ".png"
            new_url = save_base64_image(b64, ext, temp_session_id)
            url_mapping[filename] = new_url

    test_data["questions"] = _parse_questions(questions_raw, url_mapping)

    teacher = db.query(models.Teacher).filter(models.Teacher.id == teacher_id).first()
    if teacher:
        if teacher.subject:
            test_data["subject"] = teacher.subject
        if teacher.classes and not test_data.get("class_name"):
            classes_list = [c.strip() for c in teacher.classes.split(",") if c.strip()]
            if classes_list:
                test_data["class_name"] = classes_list[0]

    test = crud.create_test(db, teacher_id=teacher_id, data=test_data)
    test.folder_name = generate_new_folder_name(test)
    db.commit()
    
    if temp_session_id:
        move_temp_images_to_test(test, temp_session_id, db)
        
    save_test_locally(test)
    return test


def _parse_questions(questions_raw: list, url_mapping: dict | None = None) -> list:
    if url_mapping is None:
        url_mapping = {}
        
    allowed_q = {
        "question_text", "question_type", "points", "topic",
        "difficulty", "explanation", "order_index", "image_url",
    }
    result = []
    for i, q in enumerate(questions_raw):
        q_data = {k: v for k, v in q.items() if k in allowed_q}
        q_data.setdefault("order_index", i)
        q_data.setdefault("points", 1.0)
        q_data.setdefault("question_type", "single_choice")
        
        if q_data.get("image_url"):
            fname = os.path.basename(q_data["image_url"])
            if fname in url_mapping:
                q_data["image_url"] = url_mapping[fname]

        options_raw = q.get("options", [])
        q_data["options"] = _parse_options(options_raw, url_mapping)
        result.append(q_data)
    return result


def _parse_options(options_raw: list, url_mapping: dict | None = None) -> list:
    if url_mapping is None:
        url_mapping = {}
        
    allowed_o = {"option_text", "is_correct", "order_index", "matching_text", "image_url"}
    result = []
    for i, o in enumerate(options_raw):
        o_data = {k: v for k, v in o.items() if k in allowed_o}
        o_data.setdefault("order_index", i)
        o_data.setdefault("is_correct", False)
        
        if o_data.get("image_url"):
            fname = os.path.basename(o_data["image_url"])
            if fname in url_mapping:
                o_data["image_url"] = url_mapping[fname]
                
        result.append(o_data)
    return result


# ---------------------------------------------------------------------------
# CSV Export
# ---------------------------------------------------------------------------

def export_results_to_csv(attempts: list[models.StudentAttempt]) -> str:
    """
    Генерує CSV зі списку спроб.
    Повертає рядок CSV.
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Учень",
        "Клас",
        "Тест",
        "Бали",
        "Максимум балів",
        "Відсоток",
        "Оцінка",
        "Початок",
        "Завершення",
        "Статус",
    ])
    for attempt in attempts:
        test_title = ""
        class_name = ""
        if attempt.session and attempt.session.test:
            test_title = attempt.session.test.title
            class_name = attempt.session.test.class_name or ""

        score = attempt.score if attempt.score is not None else 0
        max_score = attempt.max_score if attempt.max_score is not None else 0
        percent = f"{score / max_score * 100:.1f}%" if max_score else "—"
        
        grade = "—"
        if max_score > 0 and attempt.session and attempt.session.test and attempt.session.test.max_grade:
            max_grade = attempt.session.test.max_grade
            grade = round((score / max_score) * max_grade)

        writer.writerow([
            attempt.student_name,
            class_name,
            test_title,
            score,
            max_score,
            percent,
            grade,
            attempt.started_at.strftime("%Y-%m-%d %H:%M") if attempt.started_at else "",
            attempt.finished_at.strftime("%Y-%m-%d %H:%M") if attempt.finished_at else "",
            attempt.status.value,
        ])
    return output.getvalue()

# ---------------------------------------------------------------------------
# XML (MyTestX) Import
# ---------------------------------------------------------------------------

def save_base64_image(base64_str: str, ext: str = ".bmp", temp_session_id: str | None = None) -> str:
    base64_str = "".join(base64_str.split())
    raw_bytes = base64.b64decode(base64_str)
    try:
        # Виправлення бага MyTestX: бінарні дані BMP були прочитані як рядок cp1251
        # і збережені в XML як UTF-8, а вже потім закодовані у Base64.
        text = raw_bytes.decode('utf-8')
        res = bytearray()
        for ch in text:
            try:
                res.extend(ch.encode('cp1251'))
            except UnicodeEncodeError:
                res.append(ord(ch) & 0xFF)
        image_data = bytes(res)
    except UnicodeDecodeError:
        image_data = raw_bytes

    filename = f"{uuid.uuid4().hex}{ext}"
    if temp_session_id:
        if not re.match(r"^[a-zA-Z0-9_\-]+$", temp_session_id):
            raise ValueError("Invalid temp session ID")
        upload_dir = os.path.join("app", "static", "tests", "temp", temp_session_id)
        os.makedirs(upload_dir, exist_ok=True)
        filepath = os.path.join(upload_dir, filename)
        url_path = f"/static/tests/temp/{temp_session_id}/{filename}"
    else:
        upload_dir = os.path.join("app", "static", "uploads")
        os.makedirs(upload_dir, exist_ok=True)
        filepath = os.path.join(upload_dir, filename)
        url_path = f"/static/uploads/{filename}"

    with open(filepath, "wb") as f:
        f.write(image_data)
    return url_path

def import_test_from_mytestx_xml(
    db: Session,
    teacher_id: int,
    xml_content: str,
    temp_session_id: str | None = None,
    fallback_title: str = "Імпортований тест"
) -> models.Test:
    try:
        root = ET.fromstring(xml_content)
    except ET.ParseError as exc:
        raise ValueError(f"Невірний XML: {exc}")

    if root.tag != "MyTestX":
        raise ValueError("Файл не є валідним форматом MyTestX (очікується тег <MyTestX>)")

    test_title = fallback_title
    questions_data = []

    # --- Priority 1: <MyTestX><Title>...
    root_title_elem = root.find("Title")
    if root_title_elem is not None and root_title_elem.text and root_title_elem.text.strip():
        test_title = root_title_elem.text.strip()

    # --- Priority 2: <MyTestX><Test Name="..."> attribute or <Test><Title>
    if test_title == fallback_title:
        test_elem = root.find("Test")
        if test_elem is not None:
            name_attr = test_elem.get("Name", "").strip()
            if name_attr:
                test_title = name_attr
            else:
                test_title_elem = test_elem.find("Title")
                if test_title_elem is not None and test_title_elem.text and test_title_elem.text.strip():
                    test_title = test_title_elem.text.strip()

    groups = root.find("Groups")
    if groups is not None:
        for group in groups.findall("Group"):
            # --- Priority 3: <Groups><Group><Title> (use only first non-empty group title if still using fallback)
            title_elem = group.find("Title")
            group_title = None
            if title_elem is not None and title_elem.text:
                group_title = title_elem.text.strip()
                if test_title == fallback_title:
                    test_title = group_title

            tasks = group.find("Tasks")
            if tasks is not None:
                for task in tasks.findall("Task"):
                    q_type_raw = task.get("Type", "TYPE_TASK_CHOICE_SINGLE")
                    score = float(task.get("Score", "1"))
                    
                    q_text = ""
                    q_text_elem = task.find("QuestionText")
                    if q_text_elem is not None:
                        pt_elem = q_text_elem.find("PlainText")
                        if pt_elem is not None and pt_elem.text:
                            q_text = pt_elem.text.strip()
                    
                    if not q_text:
                        continue

                    # Extract question image if any
                    q_image_url = None
                    q_image = task.find("QuestionImage")
                    if q_image is not None and q_image.text:
                        ext = ".bmp"
                        fname = q_image.get("FileName", "").lower()
                        if fname.endswith(".jpg"): ext = ".jpg"
                        elif fname.endswith(".png"): ext = ".png"
                        q_image_url = save_base64_image(q_image.text, ext, temp_session_id)

                    # Map question type
                    q_type = "single_choice"
                    if "MULTIPLE" in q_type_raw:
                        q_type = "multiple_choice"
                    elif "TRUE_FALSE" in q_type_raw:
                        q_type = "true_false"
                    elif "IMAGE_POINT" in q_type_raw or "PLACE_ON_IMAGE" in q_type_raw:
                        q_type = "hotspot"
                    elif "ORDER" in q_type_raw:
                        q_type = "sequence"
                    elif "COLLATION" in q_type_raw or "MATCHING" in q_type_raw:
                        q_type = "matching"
                    elif "MANUAL" in q_type_raw or "STRING" in q_type_raw or "ENTER_TEXT" in q_type_raw:
                        q_type = "short_text"

                    options_data = []
                    
                    if q_type == "hotspot":
                        # For hotspot, we extract regions
                        regions = task.find("Regions")
                        if regions is not None:
                            for r_elem in regions.findall("Region"):
                                if r_elem.text:
                                    options_data.append({
                                        "option_text": r_elem.text.strip(),
                                        "is_correct": True,
                                        "order_index": len(options_data)
                                    })
                    elif q_type == "matching":
                        # For matching, we first parse Variants2
                        variants2_list = []
                        variants2_elem = task.find("Variants2")
                        if variants2_elem is not None:
                            for v2 in variants2_elem.findall("VariantText"):
                                v2_pt = v2.find("PlainText")
                                if v2_pt is not None and v2_pt.text:
                                    variants2_list.append(v2_pt.text.strip())
                                else:
                                    variants2_list.append("")
                        
                        variants = task.find("Variants")
                        if variants is not None:
                            for i, var in enumerate(variants.findall("VariantText")):
                                opt_text = ""
                                v_pt = var.find("PlainText")
                                if v_pt is not None and v_pt.text:
                                    opt_text = v_pt.text.strip()
                                
                                # Check for variant image
                                v_image_url = None
                                v_image = var.find("VariantImage")
                                if v_image is not None and v_image.text:
                                    ext = ".bmp"
                                    fname = v_image.get("FileName", "").lower()
                                    if fname.endswith(".jpg"): ext = ".jpg"
                                    elif fname.endswith(".png"): ext = ".png"
                                    v_image_url = save_base64_image(v_image.text, ext, temp_session_id)
                                
                                correct_ans_idx_str = var.get("CorrectAnswer")
                                matching_text = ""
                                try:
                                    idx = int(correct_ans_idx_str) - 1
                                    if 0 <= idx < len(variants2_list):
                                        matching_text = variants2_list[idx]
                                except (ValueError, TypeError):
                                    pass
                                
                                options_data.append({
                                    "option_text": opt_text or f"Варіант {i+1}",
                                    "matching_text": matching_text,
                                    "is_correct": True,
                                    "order_index": i,
                                    "image_url": v_image_url
                                })
                    else:
                        variants = task.find("Variants")
                        if variants is not None:
                            for i, var in enumerate(variants.findall("VariantText")):
                                if q_type == "sequence":
                                    is_correct = True
                                    try:
                                        order_index = int(var.get("CorrectAnswer", "1")) - 1
                                    except (ValueError, TypeError):
                                        order_index = i
                                else:
                                    is_correct = var.get("CorrectAnswer", "False").lower() == "true"
                                    order_index = i
                                
                                opt_text = ""
                                v_pt = var.find("PlainText")
                                if v_pt is not None and v_pt.text:
                                    opt_text = v_pt.text.strip()
                                
                                # Check for variant image
                                v_image_url = None
                                v_image = var.find("VariantImage")
                                if v_image is not None and v_image.text:
                                    ext = ".bmp"
                                    fname = v_image.get("FileName", "").lower()
                                    if fname.endswith(".jpg"): ext = ".jpg"
                                    elif fname.endswith(".png"): ext = ".png"
                                    v_image_url = save_base64_image(v_image.text, ext, temp_session_id)
                                
                                options_data.append({
                                    "option_text": opt_text or f"Варіант {i+1}",
                                    "is_correct": is_correct,
                                    "order_index": order_index,
                                    "image_url": v_image_url
                                })
                    
                    if options_data:
                        questions_data.append({
                            "question_text": q_text,
                            "question_type": q_type,
                            "points": score,
                            "topic": group_title,
                            "order_index": len(questions_data),
                            "image_url": q_image_url,
                            "options": options_data
                        })
    
    if not questions_data:
        raise ValueError("У файлі не знайдено жодного підтримуваного питання (з варіантами відповідей)")
 
    teacher = db.query(models.Teacher).filter(models.Teacher.id == teacher_id).first()
    
    # Parse subject & class & description from XML if present
    subject_val = None
    xml_subj = root.find("Subject")
    if xml_subj is not None and xml_subj.text:
        subject_val = xml_subj.text.strip()
    else:
        test_elem = root.find("Test")
        if test_elem is not None:
            xml_subj2 = test_elem.find("Subject")
            if xml_subj2 is not None and xml_subj2.text:
                subject_val = xml_subj2.text.strip()
    if not subject_val:
        subject_val = teacher.subject if (teacher and teacher.subject) else ""

    class_val = None
    xml_class = root.find("Class")
    if xml_class is not None and xml_class.text:
        class_val = xml_class.text.strip()
    else:
        xml_class_name = root.find("ClassName")
        if xml_class_name is not None and xml_class_name.text:
            class_val = xml_class_name.text.strip()
        else:
            test_elem = root.find("Test")
            if test_elem is not None:
                xml_class2 = test_elem.find("Class")
                if xml_class2 is not None and xml_class2.text:
                    class_val = xml_class2.text.strip()
                else:
                    xml_class_name2 = test_elem.find("ClassName")
                    if xml_class_name2 is not None and xml_class_name2.text:
                        class_val = xml_class_name2.text.strip()
    if not class_val:
        if teacher and teacher.classes:
            classes_list = [c.strip() for c in teacher.classes.split(",") if c.strip()]
            if classes_list:
                class_val = classes_list[0]
        else:
            class_val = ""

    description_val = None
    xml_desc = root.find("Description")
    if xml_desc is not None and xml_desc.text:
        description_val = xml_desc.text.strip()
    else:
        test_elem = root.find("Test")
        if test_elem is not None:
            xml_desc2 = test_elem.find("Description")
            if xml_desc2 is not None and xml_desc2.text:
                description_val = xml_desc2.text.strip()
    if not description_val:
        description_val = "Імпортовано з MyTestX"
 
    test_data = {
        "title": test_title,
        "subject": subject_val,
        "class_name": class_val,
        "description": description_val,
        "time_limit_minutes": 45,
        "shuffle_questions": True,
        "shuffle_options": True,
        "show_result_after_finish": True,
        "show_correct_answers": True,
        "is_formative": False,
        "allow_partial_grading": False,
        "excluded_topics": "[]",
        "questions": questions_data
    }

    test = crud.create_test(db, teacher_id=teacher_id, data=test_data)
    test.folder_name = generate_new_folder_name(test)
    db.commit()
    
    if temp_session_id:
        move_temp_images_to_test(test, temp_session_id, db)
    save_test_locally(test)
    
    # Attach folder_name to the test object as a transient attribute for the caller
    test._folder_name = test.folder_name
    
    return test


# ---------------------------------------------------------------------------
# Word (.docx) Import
# ---------------------------------------------------------------------------

def is_run_green(run) -> bool:
    # 1. Check highlight color
    try:
        if WD_COLOR_INDEX is not None:
            hc = run.font.highlight_color
            if hc in (WD_COLOR_INDEX.GREEN, WD_COLOR_INDEX.BRIGHT_GREEN):
                return True
    except (AttributeError, ValueError, TypeError, KeyError):
        pass
        
    # 2. Check font color
    try:
        if run.font.color and run.font.color.rgb:
            color = run.font.color.rgb
            r_val = color[0]
            g_val = color[1]
            b_val = color[2]
            if g_val > 100 and g_val > r_val * 1.2 and g_val > b_val * 1.2:
                return True
    except (AttributeError, ValueError, TypeError, KeyError):
        pass
        
    # 3. Check shading/fill in XML properties (w:shd w:fill)
    try:
        rPr = getattr(run._r, 'rPr', None)
        if rPr is not None:
            shd = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
            if shd is not None:
                fill = shd.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill')
                if fill:
                    r_val = int(fill[0:2], 16)
                    g_val = int(fill[2:4], 16)
                    b_val = int(fill[4:6], 16)
                    if g_val > 100 and g_val > r_val * 1.2 and g_val > b_val * 1.2:
                        return True
    except (AttributeError, ValueError, TypeError, KeyError):
        pass
    return False


def extract_images_from_paragraph(doc, para, temp_session_id) -> list:
    image_urls = []
    namespaces = {
        'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
        'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    }
    # Find drawings in paragraph XML (both inline and anchor)
    drawings = []
    if para._element is not None:
        drawings = para._element.findall('.//wp:inline', namespaces) + para._element.findall('.//wp:anchor', namespaces)
        
    for drawing in drawings:
        pic = drawing.find('.//pic:pic', namespaces)
        if pic is not None:
            blip = pic.find('.//a:blip', namespaces)
            if blip is not None:
                rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if rId and rId in doc.part.related_parts:
                    image_part = doc.part.related_parts[rId]
                    image_bytes = image_part._blob
                    
                    ext = ".png"
                    partname = image_part.partname.lower()
                    if partname.endswith((".jpg", ".jpeg")):
                        ext = ".jpg"
                    elif partname.endswith(".gif"):
                        ext = ".gif"
                    elif partname.endswith(".bmp"):
                        ext = ".bmp"
                         
                    filename = f"{uuid.uuid4().hex}{ext}"
                    upload_dir = os.path.join("app", "static", "tests", "temp", temp_session_id)
                    os.makedirs(upload_dir, exist_ok=True)
                    filepath = os.path.join(upload_dir, filename)
                    
                    with open(filepath, "wb") as f:
                        f.write(image_bytes)
                        
                    url_path = f"/static/tests/temp/{temp_session_id}/{filename}"
                    image_urls.append(url_path)
    return image_urls


def normalize_text_for_search(t: str) -> str:
    t = t.lower()
    # Normalize Latin lookalikes to Cyrillic
    trans = {
        'i': 'і', 'a': 'а', 'b': 'в', 'c': 'с', 'e': 'е', 'h': 'н',
        'k': 'к', 'm': 'м', 'o': 'о', 'p': 'р', 't': 'т', 'x': 'х',
        'y': 'у', 's': 'ѕ'
    }
    for lat, cyr in trans.items():
        t = t.replace(lat, cyr)
    return t


def match_letters(l1: str, l2: str) -> bool:
    return normalize_text_for_search(l1) == normalize_text_for_search(l2)


def find_matching_value_by_letter(d: dict, letter: str) -> str | None:
    for k, v in d.items():
        if match_letters(str(k), letter):
            return v
    return None


def finalize_question_types_and_answers(q):
    temp_left = q.pop("temp_matching_left", {})
    temp_right = q.pop("temp_matching_right", {})
    temp_seq = q.pop("temp_sequence_options", [])
    
    # 1. Sequence questions
    if q["question_type"] == "sequence":
        # If no correct order was parsed, use document order
        if not q["options"]:
            for idx, item in enumerate(temp_seq):
                q["options"].append({
                    "option_text": item["text"],
                    "is_correct": True,
                    "order_index": idx,
                    "image_url": item["images"][0] if item["images"] else None
                })
        q.pop("temp_sequence_options", None)
        
    # 2. Matching questions
    elif q["question_type"] == "matching":
        # If no correct match parsed, match sequentially
        if not q["options"]:
            for idx, (num, left_txt) in enumerate(sorted(temp_left.items())):
                sorted_keys = sorted(temp_right.keys())
                r_txt = ""
                if idx < len(sorted_keys):
                    r_txt = temp_right[sorted_keys[idx]]
                q["options"].append({
                    "option_text": left_txt,
                    "matching_text": r_txt,
                    "is_correct": True,
                    "order_index": idx
                })
        q.pop("temp_matching_left", None)
        q.pop("temp_matching_right", None)
        
    # 3. Standard choice questions
    else:
        correct_count = sum(1 for o in q["options"] if o["is_correct"])
        if correct_count > 1:
            q["question_type"] = "multiple_choice"
        elif correct_count == 1 and q["question_type"] == "multiple_choice":
            pass


def import_test_from_docx(
    db: Session,
    teacher_id: int,
    docx_content: bytes,
    temp_session_id: str | None = None,
    fallback_title: str = "Імпортований тест"
) -> models.Test:
    if Document is None:
        raise ImportError("Бібліотека python-docx не встановлена. Встановіть її через pip install python-docx")
    
    doc = Document(io.BytesIO(docx_content))
    
    test_title = fallback_title
    title_candidate = ""
    first_non_empty = None
    
    parsed_subject = None
    parsed_class = None
    parsed_description = None
    metadata_elements = set()
    
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        m_subj = re.match(r"^\s*(Предмет|Subject)\s*:\s*(.*)$", t, re.IGNORECASE)
        m_class = re.match(r"^\s*(Клас|Class)\s*:\s*(.*)$", t, re.IGNORECASE)
        m_desc = re.match(r"^\s*(Опис|Description)\s*:\s*(.*)$", t, re.IGNORECASE)
        
        if m_subj:
            parsed_subject = m_subj.group(2).strip()
            metadata_elements.add(p._element)
        elif m_class:
            parsed_class = m_class.group(2).strip()
            metadata_elements.add(p._element)
        elif m_desc:
            parsed_description = m_desc.group(2).strip()
            metadata_elements.add(p._element)
            
    for p in doc.paragraphs:
        t = p.text.strip()
        if t and p._element not in metadata_elements:
            first_non_empty = p
            break
            
    if first_non_empty:
        t_text = first_non_empty.text.strip()
        t_text_first_line = t_text.split('\n')[0].strip()
        
        q_match = re.match(r"^\s*(Запитання|Питання)\s+(\d+)", t_text_first_line, re.IGNORECASE)
        q_num_match = re.match(r"^\s*(\d+)[\.\)]", t_text_first_line)
        if not q_match and not q_num_match:
            test_title = t_text_first_line[:120]
            title_candidate = t_text
            
    # Pre-process paragraphs into logical lines (split by \n)
    lines = []
    for para in doc.paragraphs:
        if para._element in metadata_elements:
            continue
        text = para.text
        images = extract_images_from_paragraph(doc, para, temp_session_id)
        
        parts = text.split('\n')
        for idx, part in enumerate(parts):
            part_strip = part.strip()
            if not part_strip and not (idx == 0 and images):
                continue
                
            is_bold = False
            is_green = False
            for run in para.runs:
                run_text = run.text.strip()
                if run_text and run_text in part_strip:
                    if run.bold:
                        is_bold = True
                    if is_run_green(run):
                        is_green = True
                        
            lines.append({
                "text": part_strip,
                "is_bold": is_bold,
                "is_green": is_green,
                "images": images if idx == 0 else []
            })
            
    questions_data = []
    current_q = None
    state = "IDLE"
    
    for line in lines:
        text = line["text"]
        if text == title_candidate and title_candidate:
            title_candidate = ""
            continue
            
        matching_row_match = re.match(r"^\s*(\d+)[\.\)]\s*(.*?)\s+([a-zA-Zа-яА-ЯєієїґЄІЇҐ])[\.\)]\s*(.*)$", text)
        opt_match = re.match(r"^\s*(\+)?\s*([a-zA-Zа-яА-ЯєієїґЄІЇҐ])[\.\)]\s*(.*)$", text)
        q_header_match = re.match(r"^\s*(Запитання|Питання)\s+(\d+)[\.\)]?\s*(.*)$", text, re.IGNORECASE)
        q_num_match = re.match(r"^\s*(\d+)[\.\)]\s*(.*)$", text)
        
        is_matching_ans = re.match(r"^\s*\+?\s*Правильна відповідність:\s*(.*)$", text, re.IGNORECASE)
        is_sequence_ans = re.match(r"^\s*\+?\s*Правильна послідовність:\s*(.*)$", text, re.IGNORECASE)
        is_hotspot_ans = re.match(r"^\s*\+?\s*Правильна область:\s*(.*)$", text, re.IGNORECASE)
        
        new_q_start = False
        q_text = ""
        q_type_hint = "single_choice"
        
        if q_header_match:
            new_q_start = True
            header_text = text
            q_text = q_header_match.group(3).strip()
            header_lower = normalize_text_for_search(header_text)
            if "послідовн" in header_lower or "sequence" in header_lower:
                q_type_hint = "sequence"
            elif "відповідн" in header_lower or "matching" in header_lower:
                q_type_hint = "matching"
            elif "вказівк" in header_lower or "hotspot" in header_lower:
                q_type_hint = "hotspot"
            elif "вибір зображен" in header_lower or "image_choice" in header_lower:
                q_type_hint = "image_choice"
            elif "декілька" in header_lower or "кілька" in header_lower or "multiple" in header_lower:
                q_type_hint = "multiple_choice"
            elif "коротка" in header_lower or "short" in header_lower or "введіть" in header_lower:
                q_type_hint = "short_text"
                
        elif q_num_match:
            is_matching_context = current_q and current_q["question_type"] == "matching"
            if not is_matching_context and not matching_row_match:
                new_q_start = True
                q_text = q_num_match.group(2).strip()
                
        if new_q_start:
            if current_q:
                finalize_question_types_and_answers(current_q)
                questions_data.append(current_q)
                
            current_q = {
                "question_text": q_text,
                "question_type": q_type_hint,
                "points": 1.0,
                "topic": None,
                "order_index": len(questions_data),
                "image_url": line["images"][0] if line["images"] else None,
                "options": [],
                "temp_matching_left": {},
                "temp_matching_right": {},
                "temp_sequence_options": []
            }
            state = "IN_QUESTION"
            continue
            
        if not current_q:
            continue
            
        if is_matching_ans:
            pairs_str = re.findall(r"(\d+)\s*-\s*([a-zA-Zа-яА-ЯєієїґЄІЇҐ])", is_matching_ans.group(1))
            for pair in pairs_str:
                num = int(pair[0])
                letter = pair[1]
                left_text = find_matching_value_by_letter(current_q["temp_matching_left"], str(num)) or f"Елемент {num}"
                right_text = find_matching_value_by_letter(current_q["temp_matching_right"], letter) or f"Варіант {letter}"
                
                current_q["options"].append({
                    "option_text": left_text,
                    "matching_text": right_text,
                    "is_correct": True,
                    "order_index": len(current_q["options"])
                })
            state = "IN_OPTIONS"
            
        elif is_sequence_ans:
            seq_letters = re.findall(r"([a-zA-Zа-яА-ЯєієїґЄІЇҐ])", is_sequence_ans.group(1))
            for idx, letter in enumerate(seq_letters):
                item = None
                for candidate in current_q["temp_sequence_options"]:
                    if match_letters(candidate["letter"], letter):
                        item = candidate
                        break
                if item:
                    current_q["options"].append({
                        "option_text": item["text"],
                        "is_correct": True,
                        "order_index": idx,
                        "image_url": item["images"][0] if item["images"] else None
                    })
            state = "IN_OPTIONS"
            
        elif is_hotspot_ans:
            region_str = is_hotspot_ans.group(1).strip()
            current_q["options"].append({
                "option_text": region_str,
                "is_correct": True,
                "order_index": 0,
                "image_url": None
            })
            state = "IN_OPTIONS"
            
        elif matching_row_match:
            num = int(matching_row_match.group(1))
            left_text = matching_row_match.group(2).strip()
            right_letter = matching_row_match.group(3).upper()
            right_text = matching_row_match.group(4).strip()
            
            current_q["temp_matching_left"][num] = left_text
            current_q["temp_matching_right"][right_letter] = right_text
            state = "IN_QUESTION"
            
        elif opt_match:
            is_correct_marker = opt_match.group(1) is not None
            opt_letter = opt_match.group(2).upper()
            opt_text = opt_match.group(3).strip()
            
            is_correct = is_correct_marker
            if "(+)" in opt_text or "(+)" in text:
                is_correct = True
                opt_text = re.sub(r"\(\+\)", "", opt_text).strip()
            elif not is_correct:
                is_correct = line["is_bold"] or line["is_green"]
                
            if current_q["question_type"] == "matching":
                current_q["temp_matching_right"][opt_letter] = opt_text
            elif current_q["question_type"] == "sequence":
                current_q["temp_sequence_options"].append({
                    "letter": opt_letter,
                    "text": opt_text,
                    "images": line["images"]
                })
            else:
                current_q["options"].append({
                    "option_text": opt_text,
                    "is_correct": is_correct,
                    "order_index": len(current_q["options"]),
                    "image_url": line["images"][0] if line["images"] else None
                })
            state = "IN_OPTIONS"
            
        else:
            if not text and not line["images"]:
                continue
                
            if state == "IN_QUESTION" and current_q:
                if text:
                    if current_q["question_text"]:
                        current_q["question_text"] += "\n" + text
                    else:
                        current_q["question_text"] = text
                if line["images"] and not current_q["image_url"]:
                    current_q["image_url"] = line["images"][0]
            elif state == "IN_OPTIONS" and current_q:
                if current_q["question_type"] == "sequence" and current_q["temp_sequence_options"]:
                    last_opt = current_q["temp_sequence_options"][-1]
                    if text:
                        last_opt["text"] += "\n" + text
                    if line["images"] and not last_opt["images"]:
                        last_opt["images"] = line["images"]
                elif current_q["options"]:
                    last_opt = current_q["options"][-1]
                    if text:
                        last_opt["option_text"] += "\n" + text
                    if line["images"] and not last_opt["image_url"]:
                        last_opt["image_url"] = line["images"][0]
                        
    if current_q:
        finalize_question_types_and_answers(current_q)
        questions_data.append(current_q)
        
    if not questions_data:
        raise ValueError("У файлі не знайдено жодного запитання. Перевірте форматування (питання мають починатися з цифри з крапкою/дужкою, варіанти — з літери з дужкою).")
        
    teacher = db.query(models.Teacher).filter(models.Teacher.id == teacher_id).first()
    
    subject_val = parsed_subject
    if not subject_val:
        subject_val = teacher.subject if (teacher and teacher.subject) else ""
        
    class_val = parsed_class
    if not class_val:
        if teacher and teacher.classes:
            classes_list = [c.strip() for c in teacher.classes.split(",") if c.strip()]
            if classes_list:
                class_val = classes_list[0]
        else:
            class_val = ""
            
    description_val = parsed_description or "Імпортовано з Word (.docx)"
            
    test_data = {
        "title": test_title,
        "subject": subject_val,
        "class_name": class_val,
        "description": description_val,
        "time_limit_minutes": 45,
        "shuffle_questions": True,
        "shuffle_options": True,
        "show_result_after_finish": True,
        "show_correct_answers": True,
        "is_formative": False,
        "allow_partial_grading": False,
        "excluded_topics": "[]",
        "questions": questions_data
    }
    
    test = crud.create_test(db, teacher_id=teacher_id, data=test_data)
    test.folder_name = generate_new_folder_name(test)
    db.commit()
    
    if temp_session_id:
        move_temp_images_to_test(test, temp_session_id, db)
    save_test_locally(test)
    
    test._folder_name = test.folder_name
    return test
